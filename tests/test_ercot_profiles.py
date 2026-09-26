from datetime import date
from pathlib import Path

import numpy as np
import pytest
from openpyxl import Workbook

from pipeline.sources.ercot_profiles import (
    intervals_from_row,
    profile_days_many,
    read_profile_day,
    span_kwh,
    typical_day,
)


def test_intervals_drop_trailing_empty_cells_and_ignore_addtime():
    header = ["PType_WZ", "Date", "int_kWh1", "int_kWh2", "int_kWh3", "ADDTIME"]
    row = ["RESHIWR_NCENT", date(2018, 3, 11), 0.26, 0.25, None, date(2018, 3, 15)]
    assert intervals_from_row(header, row).tolist() == [0.26, 0.25]


def test_intervals_reject_a_gap_in_the_middle():
    header = ["int_kWh1", "int_kWh2", "int_kWh3"]
    with pytest.raises(ValueError, match="middle"):
        intervals_from_row(header, [1.0, None, 1.0])


def _workbook(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "June"
    sheet.append(["PType_WZ", "Date", "int_kWh1", "int_kWh2", "int_kWh3", "ADDTIME"])
    sheet.append(["RESHIWR_NCENT", date(2021, 6, 1), 0.2, 0.3, None, date(2021, 6, 5)])
    sheet.append(["RESLOWR_COAST", date(2021, 6, 1), 0.4, 0.5, 0.6, None])
    book.save(path)
    book.close()


def test_typical_day_keeps_96_interval_days_only():
    short = np.ones(92)
    low = np.arange(96, dtype=float)
    high = low + 10
    typical = typical_day([short, low, high])
    assert len(typical) == 96
    assert typical[0] == 5.0


def test_span_kwh_concatenates_consecutive_days(tmp_path: Path):
    path = tmp_path / "profiles.xlsx"
    _workbook(path)
    trace = span_kwh(path, date(2021, 6, 1), 1, "RESLOWR", "COAST")
    assert trace.tolist() == [0.4, 0.5, 0.6]


def test_duplicate_month_sheets_keep_the_fuller_one(tmp_path: Path):
    book = Workbook()
    partial = book.active
    partial.title = "September"
    full = book.create_sheet("Sep")
    header = ["PType_WZ", "Date", "int_kWh1", "ADDTIME"]
    partial.append(header)
    partial.append(["RESLOWR_COAST", date(2025, 9, 1), 9.0, None])
    full.append(header)
    full.append(["RESLOWR_COAST", date(2025, 9, 1), 1.0, None])
    full.append(["RESLOWR_COAST", date(2025, 9, 2), 2.0, None])
    path = tmp_path / "profiles.xlsx"
    book.save(path)
    book.close()
    days = profile_days_many(path, {"RESLOWR_COAST"})["RESLOWR_COAST"]
    assert [(day.day, trace.tolist()) for day, trace in days] == [(1, [1.0]), (2, [2.0])]


def test_read_profile_day_from_a_workbook(tmp_path: Path):
    path = tmp_path / "profiles.xlsx"
    _workbook(path)
    trace = read_profile_day(path, date(2021, 6, 1), "RESHIWR", "NCENT")
    assert trace.tolist() == pytest.approx([0.2, 0.3])
    other = read_profile_day(path, date(2021, 6, 1), "RESLOWR", "COAST")
    assert np.array_equal(other, np.array([0.4, 0.5, 0.6]))
