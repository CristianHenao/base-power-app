"""Read one day from an ERCOT backcasted load-profile workbook.

Each yearly zip holds one workbook. Months are sheets. A row is one profile
on one day, and int_kWh1... are 15-minute kWh. Short and long days leave the
extra columns empty. ADDTIME is not an interval.
"""
from __future__ import annotations

import io
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from pipeline import settings

_MONTHS = ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec")


def profile_code(profile_type: str, weather_zone: str) -> str:
    return f"{profile_type}_{weather_zone}"


def yearly_zip(year: int, directory: Path = settings.RAW_ERCOT_DIR) -> Path:
    matches = sorted(path for path in directory.glob(f"*{year}*.zip") if path.is_file())
    if len(matches) != 1:
        raise FileNotFoundError(f"expected one {year} load-profile zip in {directory}, found {len(matches)}")
    return matches[0]


def intervals_from_row(header: list[object], row: list[object]) -> np.ndarray:
    """15-minute kWh in column order. Trailing empty cells are the short-day pad."""
    values = [
        row[i]
        for i, name in enumerate(header)
        if str(name).lower().startswith("int_kwh")
    ]
    while values and values[-1] is None:
        values.pop()
    if any(value is None for value in values):
        raise ValueError("an interval is missing in the middle of the day")
    return np.asarray(values, dtype=float)


def _sheet_name(names: list[str], month: int) -> str:
    prefix = _MONTHS[month - 1]
    matches = [name for name in names if name.lower().startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"expected one sheet for month {month}, found {matches}")
    return matches[0]


def _workbook(source: Path):
    if source.suffix.lower() == ".zip":
        with zipfile.ZipFile(source) as archive:
            name = next(item for item in archive.namelist() if item.lower().endswith(".xlsx"))
            data = io.BytesIO(archive.read(name))
        return load_workbook(data, read_only=True, data_only=True)
    return load_workbook(source, read_only=True, data_only=True)


def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def typical_day(traces: list[np.ndarray], intervals: int = 96) -> np.ndarray:
    """Median 15-minute shape. Short and long daylight-saving days are left out."""
    regular = [np.asarray(trace, float) for trace in traces if len(trace) == intervals]
    if not regular:
        raise ValueError(f"no {intervals}-interval days")
    return np.median(np.stack(regular), axis=0)


def annual_kwh(days_by_month: dict[int, list[tuple[date, np.ndarray]]]) -> float:
    return float(sum(trace.sum() for days in days_by_month.values() for _, trace in days))


def _days_on_sheet(sheet, code: str) -> list[tuple[date, np.ndarray]]:
    rows = sheet.iter_rows(values_only=True)
    header = list(next(rows))
    found: list[tuple[date, np.ndarray]] = []
    for row in rows:
        if row[0] != code:
            continue
        day = _as_date(row[1])
        if day is None:
            continue
        found.append((day, intervals_from_row(header, list(row))))
    return found


def profile_month(source: Path, month: int, profile_type: str, weather_zone: str) -> list[tuple[date, np.ndarray]]:
    """Every day of one profile in one month, in sheet order."""
    code = profile_code(profile_type, weather_zone)
    workbook = _workbook(source)
    try:
        sheet = workbook[_sheet_name(workbook.sheetnames, month)]
        return _days_on_sheet(sheet, code)
    finally:
        workbook.close()


def profile_days_by_month(source: Path, profile_type: str, weather_zone: str) -> dict[int, list[tuple[date, np.ndarray]]]:
    """All twelve months for one profile. The workbook is opened once."""
    code = profile_code(profile_type, weather_zone)
    workbook = _workbook(source)
    days = {month: [] for month in range(1, 13)}
    try:
        for month in range(1, 13):
            sheet = workbook[_sheet_name(workbook.sheetnames, month)]
            days[month] = _days_on_sheet(sheet, code)
    finally:
        workbook.close()
    return days


def span_kwh(source: Path, start: date, n_days: int, profile_type: str, weather_zone: str) -> np.ndarray:
    """Concatenate daily kWh from `start` for `n_days`. All of those days must be in `source`."""
    if n_days < 1:
        raise ValueError("n_days must be positive")
    needed = [start + timedelta(days=offset) for offset in range(n_days)]
    months = sorted({day.month for day in needed})
    lookup: dict[date, np.ndarray] = {}
    for month in months:
        for day, trace in profile_month(source, month, profile_type, weather_zone):
            lookup[day] = trace
    missing = [day.isoformat() for day in needed if day not in lookup]
    if missing:
        raise ValueError(f"missing profile days: {', '.join(missing)}")
    return np.concatenate([lookup[day] for day in needed])


def read_profile_day(
    source: Path,
    day: date,
    profile_type: str,
    weather_zone: str,
) -> np.ndarray:
    """kWh per 15 minutes for one profile on one day."""
    code = profile_code(profile_type, weather_zone)
    workbook = _workbook(source)
    try:
        sheet = workbook[_sheet_name(workbook.sheetnames, day.month)]
        rows = sheet.iter_rows(values_only=True)
        header = list(next(rows))
        found: np.ndarray | None = None
        for row in rows:
            if row[0] != code or _as_date(row[1]) != day:
                continue
            if found is not None:
                raise ValueError(f"more than one {code} row on {day.isoformat()}")
            found = intervals_from_row(header, list(row))
    finally:
        workbook.close()
    if found is None:
        raise ValueError(f"no {code} row on {day.isoformat()}")
    return found
